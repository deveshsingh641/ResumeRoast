import docx
import pptx
import os
import json

def inspect_template():
    doc = docx.Document('InternshipProjectReportFormat (1).docx')
    with open('template_dump.txt', 'w', encoding='utf-8') as f:
        f.write(f'Paragraphs ({len(doc.paragraphs)}):\n')
        for i, p in enumerate(doc.paragraphs):
            runs_info = []
            for r in p.runs:
                font_name = r.font.name
                font_size = r.font.size.pt if r.font.size else None
                bold = r.bold
                italic = r.italic
                runs_info.append(f"['{r.text}' fn={font_name} sz={font_size} b={bold} i={italic}]")
            f.write(f'P{i} [{p.style.name}] align={p.alignment}: {p.text}\n')
            if runs_info:
                f.write(f'    Runs: {" ".join(runs_info)}\n')
        
        f.write(f'\nTables ({len(doc.tables)}):\n')
        for t_idx, t in enumerate(doc.tables):
            f.write(f'\n--- Table {t_idx} ({len(t.rows)}x{len(t.columns)}) ---\n')
            for r_idx, r in enumerate(t.rows):
                row_txt = [c.text.replace('\n', ' ') for c in r.cells]
                f.write(f'  Row {r_idx}: {" | ".join(row_txt)}\n')

        # Check inline shapes / images
        f.write(f'\nInline shapes: {len(doc.inline_shapes)}\n')
        for i, s in enumerate(doc.inline_shapes):
            f.write(f'Shape {i}: type={s.type}, width={s.width}, height={s.height}\n')

def inspect_pptx():
    prs = pptx.Presentation('ABES_Internship_Presentation_Devesh_Singh_Updated.pptx')
    with open('pptx_dump.txt', 'w', encoding='utf-8') as f:
        f.write(f'Total Slides: {len(prs.slides)}\n')
        for idx, slide in enumerate(prs.slides):
            f.write(f'\n=== SLIDE {idx+1} ===\n')
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            f.write(f'  [Text]: {text}\n')
                elif shape.has_table:
                    f.write(f'  [Table]: {len(shape.table.rows)}x{len(shape.table.columns)}\n')
                    for r in shape.table.rows:
                        row_vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
                        f.write(f'    {" | ".join(row_vals)}\n')
                elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
                    f.write(f'  [Picture]: name={shape.name}\n')

if __name__ == '__main__':
    inspect_template()
    inspect_pptx()
    print("Inspection complete.")
