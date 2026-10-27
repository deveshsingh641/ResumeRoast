import docx

doc = docx.Document('InternshipProjectReportFormat (1).docx')
print("Sections:")
for i, s in enumerate(doc.sections):
    print(f"Section {i}: start_type={s.start_type}, top={s.top_margin.inches}, bot={s.bottom_margin.inches}, left={s.left_margin.inches}, right={s.right_margin.inches}, orientation={s.orientation}, page_w={s.page_width.inches}, page_h={s.page_height.inches}")
    # headers / footers
    print(f"  Header text: {[p.text for p in s.header.paragraphs if p.text.strip()]}")
    print(f"  Footer text: {[p.text for p in s.footer.paragraphs if p.text.strip()]}")

print("\nTable 0 (Title page college header):")
t0 = doc.tables[0]
for c_idx, cell in enumerate(t0.rows[0].cells):
    print(f"Cell {c_idx}: width={cell.width}, text={repr(cell.text.strip())}, has_drawing={'drawing' in cell._element.xml}")

print("\nTable 1 (Signatures):")
t1 = doc.tables[1]
for c_idx, cell in enumerate(t1.rows[0].cells):
    print(f"Cell {c_idx}: text={repr(cell.text.strip())}")

print("\nTable 2 (TOC table):")
t2 = doc.tables[2]
for r_idx, r in enumerate(t2.rows):
    print(f"Row {r_idx:2d}: {[c.text.strip() for c in r.cells]}")

print("\nTable 3 (Formatting Standards):")
t3 = doc.tables[3]
for r_idx, r in enumerate(t3.rows):
    print(f"Row {r_idx:2d}: {[c.text.strip() for c in r.cells]}")
