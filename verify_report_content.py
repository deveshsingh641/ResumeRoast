import docx

doc = docx.Document('Devesh_Singh_Internship_Report_Final.docx')
found_mini = []
found_internship = []

def check_text(text, location):
    if 'mini project' in text.lower():
        found_mini.append((location, text.strip()))
    if 'internship' in text.lower():
        found_internship.append((location, text.strip()))

for i, p in enumerate(doc.paragraphs):
    check_text(p.text, f'Paragraph {i}')

for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            for p_idx, p in enumerate(cell.paragraphs):
                check_text(p.text, f'Table {t_idx} Row {r_idx} Col {c_idx}')

for s_idx, section in enumerate(doc.sections):
    for h in [section.header, section.first_page_header]:
        if h:
            for p in h.paragraphs:
                check_text(p.text, f'Section {s_idx} Header')
    for f in [section.footer, section.first_page_footer]:
        if f:
            for p in f.paragraphs:
                check_text(p.text, f'Section {s_idx} Footer')

print(f"Total occurrences of 'mini project': {len(found_mini)}")
for loc, txt in found_mini:
    print(f"  FOUND in {loc}: {txt}")

print(f"Total occurrences of 'internship': {len(found_internship)}")
print("Sample occurrences of internship:")
for loc, txt in found_internship[:8]:
    print(f"  {loc}: {txt[:80]}")
