import sys
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

deck_path = r'd:\ResumeRoast\ABES_Internship_Presentation_Devesh_Singh_Updated.pptx'
prs = pptx.Presentation(deck_path)

print(f"Total Slides: {len(prs.slides)}")
assert len(prs.slides) == 11, f"Expected 11 slides, got {len(prs.slides)}"

forbidden = ['github', 'http', 'www', '.in', '.com', 'resumeroast', 'resume roast', 'ce/it', 'ce / it']
violations = []

for idx, slide in enumerate(prs.slides):
    slide_num = idx + 1
    text_content = []
    images = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                t = p.text.strip()
                if t:
                    text_content.append(t)
        if shape.has_table:
            for row in shape.table.rows:
                for c in row.cells:
                    ct = c.text.strip()
                    if ct:
                        text_content.append(ct)
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            images.append(shape)
    
    full_text = ' '.join(text_content)
    lower = full_text.lower()
    for f in forbidden:
        if f in lower:
            # Allow legitimate academic term 'http status code'
            if f == 'http' and 'http status code' in lower and lower.count('http') == lower.count('http status code'):
                continue
            violations.append((slide_num, f, full_text[:120]))
    
    title = text_content[0] if text_content else "No Title"
    print(f"Slide {slide_num:2d}: {title:<45} | Shapes: {len(slide.shapes):2d} | Images: {len(images)}")
    if slide_num == 6:
        print(f"   --> Slide 6 Images count: {len(images)}")
        for img_idx, img in enumerate(images):
            w_in = img.width / 914400
            h_in = img.height / 914400
            print(f"       Image {img_idx+1}: ({w_in:.2f} in x {h_in:.2f} in, ratio={w_in/h_in:.2f})")
    elif slide_num == 10:
        print(f"   --> Slide 10 (Certificate) Images count: {len(images)}")
        for img_idx, img in enumerate(images):
            w_in = img.width / 914400
            h_in = img.height / 914400
            print(f"       Certificate Image: ({w_in:.2f} in x {h_in:.2f} in, ratio={w_in/h_in:.2f})")

if violations:
    print("\nVIOLATIONS FOUND:")
    for v in violations:
        print(f"  Slide {v[0]}: matched '{v[1]}' in snippet: {v[2]}")
else:
    print("\nSUCCESS: All 11 slides are 100% compliant, cleanly structured, and free of unauthorized URLs/text!")


